"""Shared helpers for building dashboard chart series and exporting
them to Excel (data + native Excel line chart). Used by the Manage page."""
from __future__ import annotations
import io

import pandas as pd
import streamlit as st
import openpyxl
from openpyxl.chart import LineChart, Reference

from utils import sheets_db, data_fetch, returns


# ----------------------------------------------------------------- loaders --
def load_holdings(tab_name: str) -> list[dict]:
    df = sheets_db.read_df(tab_name)
    holdings = []
    for _, row in df.iterrows():
        try:
            holdings.append({
                "ticker": str(row["ticker"]).strip(),
                "asset_type": str(row.get("asset_type", "stock")).strip().lower() or "stock",
                "weight": float(row["weight"]) / 100.0,
                "inception_date": pd.to_datetime(row["purchase_date"]),
            })
        except (KeyError, ValueError, TypeError):
            continue
    return holdings


def load_backtest(portfolio_label: str) -> pd.Series:
    df = sheets_db.read_df("backtest_history")
    if df.empty:
        return pd.Series(dtype=float)
    df = df[df["portfolio"] == portfolio_label].copy()
    if df.empty:
        return pd.Series(dtype=float)
    df["date"] = pd.to_datetime(df["date"])
    df = df.sort_values("date")
    return pd.Series(pd.to_numeric(df["index_value"], errors="coerce").values, index=df["date"])


# ----------------------------------------------- cached portfolio index build --
@st.cache_data(ttl=3600, show_spinner=False)
def _compute_portfolio_index_cached(tab_name: str, portfolio_label: str, holdings_key: tuple) -> pd.Series:
    holdings = [
        {"ticker": t, "asset_type": at, "weight": w, "inception_date": pd.Timestamp(d)}
        for t, at, w, d in holdings_key
    ]
    price_data = data_fetch.get_prices_batch(holdings)
    backtest_index_values = load_backtest(portfolio_label)
    rebalance_freq = sheets_db.get_rebalance_frequency(portfolio_label)
    live_start_date = backtest_index_values.index[-1] if not backtest_index_values.empty else None

    live_index = returns.compute_live_index(
        holdings, price_data,
        rebalance_frequency=rebalance_freq,
        live_start_date=live_start_date,
    )
    if live_index.empty and holdings:
        live_index = returns.compute_live_index(
            holdings, price_data,
            rebalance_frequency=rebalance_freq,
            live_start_date=None,
        )
        if not live_index.empty and live_start_date is not None:
            live_index = live_index[live_index.index >= live_start_date]
    return returns.chain_link_backtest(backtest_index_values, live_index)


def compute_portfolio_index(tab_name: str, portfolio_label: str, holdings: list[dict]) -> pd.Series:
    holdings_key = tuple(
        (h["ticker"], h["asset_type"], h["weight"], pd.Timestamp(h["inception_date"]))
        for h in holdings
    )
    return _compute_portfolio_index_cached(tab_name, portfolio_label, holdings_key)


# ----------------------------------------------------------- series options --
def build_series_options() -> dict:
    """Reproduces exactly the dashboard's series_options dict."""
    portfolio_labels = sheets_db.get_portfolios()
    backtest_df = sheets_db.read_df("backtest_history")
    series_options = {}
    for tab_name, label in portfolio_labels.items():
        holdings = load_holdings(tab_name)
        if holdings or not backtest_df[backtest_df["portfolio"] == label].empty:
            series_options[label] = compute_portfolio_index(tab_name, label, holdings)
    watchlist_df = sheets_db.read_df("watchlist_etfs")
    if not watchlist_df.empty:
        watchlist_prices = data_fetch.get_watchlist_prices(watchlist_df)
        series_options.update(watchlist_prices)
    order_map = sheets_db.get_display_order()
    sorted_keys = sorted(series_options.keys(), key=lambda x: order_map.get(x, 9999))
    return {k: series_options[k] for k in sorted_keys}


# --------------------------------------------------------- period slicer --
def slice_by_period(chart_series: pd.Series, period: str) -> pd.Series:
    chart_series = chart_series.dropna()
    if chart_series.empty:
        return chart_series
    last_date = chart_series.index.max()
    if period == "5D":
        start_date = last_date - pd.Timedelta(days=5)
    elif period == "1M":
        start_date = last_date - pd.DateOffset(months=1)
    elif period == "3M":
        start_date = last_date - pd.DateOffset(months=3)
    elif period == "6M":
        start_date = last_date - pd.DateOffset(months=6)
    elif period == "YTD":
        start_date = pd.Timestamp(year=last_date.year, month=1, day=1)
    elif period == "1Y":
        start_date = last_date - pd.DateOffset(years=1)
    elif period == "3Y":
        start_date = last_date - pd.DateOffset(years=3)
    elif period == "5Y":
        start_date = last_date - pd.DateOffset(years=5)
    else:  # "Max"
        start_date = chart_series.index.min()
    return chart_series[chart_series.index >= start_date]


# ----------------------------------------------------------- Excel builder --
def build_excel_bytes(label: str, chart_series: pd.Series, period: str) -> bytes:
    """Returns an .xlsx file (as bytes) containing:
       - 'Chart Data' sheet: Date, Index Value, Total Return (%)
       - 'Chart' sheet: native Excel LineChart of Total Return (%)"""
    view = slice_by_period(chart_series, period)
    if view.empty or len(view) < 2:
        raise ValueError("Not enough data to export for the selected period.")

    rebased = (view / view.iloc[0] - 1) * 100
    df = pd.DataFrame({
        "Date": view.index,
        "Index Value": view.values,
        "Total Return (%)": rebased.values,
    })
    df["Date"] = pd.to_datetime(df["Date"]).dt.date

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chart Data"
    ws.append(["Date", "Index Value", "Total Return (%)"])
    for _, row in df.iterrows():
        ws.append([row["Date"], float(row["Index Value"]), round(float(row["Total Return (%)"]), 4)])

    # Format the date column
    for r in range(2, len(df) + 2):
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"

    # Native Excel line chart on its own sheet
    chart = LineChart()
    chart.title = f"{label} - Total Return % ({period})"
    chart.y_axis.title = "Total return (%)"
    chart.x_axis.title = "Date"
    chart.height = 12
    chart.width = 24

    data_ref = Reference(ws, min_col=3, min_row=1, max_row=len(df) + 1, max_col=3)
    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=len(df) + 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.legend = None

    ws_chart = wb.create_sheet("Chart")
    ws_chart.add_chart(chart, "A1")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
